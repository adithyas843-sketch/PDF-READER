
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font

def create_excel_report(df, exc_df):

    wb = Workbook()

    ws = wb.active
    ws.title = "PO Summary"

    for c, h in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)

    for r, row in enumerate(df.values.tolist(), start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    ws2 = wb.create_sheet("Exceptions")

    if not exc_df.empty:
        for c, h in enumerate(exc_df.columns, start=1):
            ws2.cell(row=1, column=c, value=h)

        for r, row in enumerate(exc_df.values.tolist(), start=2):
            for c, value in enumerate(row, start=1):
                ws2.cell(row=r, column=c, value=value)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()
